"""Extracteur de texte pour différents formats de documents."""

import os
from pathlib import Path
from typing import Optional
from enum import Enum


class DocumentType(Enum):
    """Types de documents supportés."""
    PDF = "pdf"
    DOCX = "docx"
    DOC = "doc"
    TXT = "txt"
    UNKNOWN = "unknown"


class DocumentExtractor:
    """
    Extracteur de texte pour documents PDF et DOCX.

    Cette classe fournit des méthodes pour extraire le texte
    de différents formats de documents.
    """

    def __init__(self, use_ocr: bool = False):
        """
        Initialise l'extracteur.

        Args:
            use_ocr: Utiliser l'OCR pour les PDF scannés
        """
        self.use_ocr = use_ocr

    @staticmethod
    def detect_type(file_path: str) -> DocumentType:
        """
        Détecte le type de document à partir de l'extension.

        Args:
            file_path: Chemin vers le fichier

        Returns:
            Type de document
        """
        ext = Path(file_path).suffix.lower().lstrip('.')

        type_mapping = {
            'pdf': DocumentType.PDF,
            'docx': DocumentType.DOCX,
            'doc': DocumentType.DOC,
            'txt': DocumentType.TXT,
        }

        return type_mapping.get(ext, DocumentType.UNKNOWN)

    def extract(self, file_path: str) -> str:
        """
        Extrait le texte d'un document.

        Args:
            file_path: Chemin vers le fichier

        Returns:
            Texte extrait du document

        Raises:
            FileNotFoundError: Si le fichier n'existe pas
            ValueError: Si le type de document n'est pas supporté
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Fichier non trouvé: {file_path}")

        doc_type = self.detect_type(file_path)

        extractors = {
            DocumentType.PDF: self._extract_pdf,
            DocumentType.DOCX: self._extract_docx,
            DocumentType.TXT: self._extract_txt,
        }

        extractor = extractors.get(doc_type)
        if not extractor:
            raise ValueError(f"Type de document non supporté: {doc_type.value}")

        return extractor(file_path)

    def _extract_pdf(self, file_path: str) -> str:
        """
        Extrait le texte d'un fichier PDF.

        Args:
            file_path: Chemin vers le fichier PDF

        Returns:
            Texte extrait
        """
        text_parts = []

        # Essayer avec pdfplumber d'abord (meilleure extraction)
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)

            if text_parts:
                return '\n\n'.join(text_parts)

        except ImportError:
            pass
        except Exception:
            pass

        # Fallback sur PyPDF2
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(file_path)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

            if text_parts:
                return '\n\n'.join(text_parts)

        except ImportError:
            pass
        except Exception:
            pass

        # Essayer OCR si activé et aucun texte trouvé
        if self.use_ocr and not text_parts:
            return self._extract_pdf_ocr(file_path)

        return '\n\n'.join(text_parts) if text_parts else ""

    def _extract_pdf_ocr(self, file_path: str) -> str:
        """
        Extrait le texte d'un PDF scanné via OCR.

        Args:
            file_path: Chemin vers le fichier PDF

        Returns:
            Texte extrait via OCR
        """
        try:
            import pytesseract
            from PIL import Image
            from pdf2image import convert_from_path

            # Convertir PDF en images
            images = convert_from_path(file_path)

            text_parts = []
            for image in images:
                text = pytesseract.image_to_string(image, lang='fra')
                if text:
                    text_parts.append(text)

            return '\n\n'.join(text_parts)

        except ImportError:
            return ""
        except Exception:
            return ""

    def _extract_docx(self, file_path: str) -> str:
        """
        Extrait le texte d'un fichier DOCX.

        Args:
            file_path: Chemin vers le fichier DOCX

        Returns:
            Texte extrait
        """
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts = []

            # Extraire les paragraphes
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Extraire les tableaux
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        text_parts.append(row_text)

            return '\n'.join(text_parts)

        except ImportError:
            raise ImportError(
                "python-docx est requis pour lire les fichiers DOCX. "
                "Installez-le avec: pip install python-docx"
            )

    def _extract_txt(self, file_path: str) -> str:
        """
        Lit un fichier texte.

        Args:
            file_path: Chemin vers le fichier texte

        Returns:
            Contenu du fichier
        """
        encodings = ['utf-8', 'latin-1', 'cp1252']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue

        # Dernier recours: ignorer les erreurs
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def extract_with_metadata(self, file_path: str) -> dict:
        """
        Extrait le texte et les métadonnées d'un document.

        Args:
            file_path: Chemin vers le fichier

        Returns:
            Dictionnaire avec le texte et les métadonnées
        """
        text = self.extract(file_path)
        doc_type = self.detect_type(file_path)

        metadata = {
            'file_path': file_path,
            'file_name': os.path.basename(file_path),
            'file_size': os.path.getsize(file_path),
            'document_type': doc_type.value,
            'text_length': len(text),
            'text': text,
        }

        # Extraire les métadonnées PDF si disponible
        if doc_type == DocumentType.PDF:
            pdf_meta = self._get_pdf_metadata(file_path)
            if pdf_meta:
                metadata['pdf_metadata'] = pdf_meta

        return metadata

    def _get_pdf_metadata(self, file_path: str) -> Optional[dict]:
        """
        Extrait les métadonnées d'un fichier PDF.

        Args:
            file_path: Chemin vers le fichier PDF

        Returns:
            Métadonnées du PDF ou None
        """
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(file_path)
            meta = reader.metadata

            if meta:
                return {
                    'title': meta.title,
                    'author': meta.author,
                    'subject': meta.subject,
                    'creator': meta.creator,
                    'producer': meta.producer,
                    'creation_date': str(meta.creation_date) if meta.creation_date else None,
                    'modification_date': str(meta.modification_date) if meta.modification_date else None,
                    'num_pages': len(reader.pages),
                }

        except Exception:
            pass

        return None
